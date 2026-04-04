import hashlib
import os
import re
import subprocess
import uuid
import zipfile
from html import escape
from xml.etree import ElementTree as ET
from datetime import datetime
from pathlib import Path
from django.conf import settings


def get_channel_upload_path(workspace_id, channel_id, filename):
    now = datetime.now()
    safe_name = f"{uuid.uuid4().hex[:12]}_{filename}"
    return Path(
        settings.STORAGE_ROOT,
        'workspaces', str(workspace_id),
        'channels', str(channel_id),
        str(now.year), f'{now.month:02d}',
        safe_name,
    )


def get_dm_upload_path(user_id, thread_id, filename):
    now = datetime.now()
    safe_name = f"{uuid.uuid4().hex[:12]}_{filename}"
    return Path(
        settings.STORAGE_ROOT,
        'users', str(user_id),
        'dm', str(thread_id),
        str(now.year), f'{now.month:02d}',
        safe_name,
    )


def get_avatar_path(user_id, filename):
    ext = filename.rsplit('.', 1)[-1] if '.' in filename else 'png'
    return Path(
        settings.STORAGE_ROOT,
        'avatars',
        f'{user_id}_{uuid.uuid4().hex[:8]}.{ext}',
    )


def save_uploaded_file(file_obj, dest_path):
    """Save an uploaded file and return its checksum."""
    dest_path = Path(dest_path)
    dest_path.parent.mkdir(parents=True, exist_ok=True)

    hasher = hashlib.sha256()
    with open(dest_path, 'wb') as dest:
        for chunk in file_obj.chunks():
            dest.write(chunk)
            hasher.update(chunk)

    return hasher.hexdigest()


def delete_file(path):
    """Delete a file from storage."""
    try:
        os.remove(path)
    except OSError:
        pass


OFFICE_MIME_TYPES = {
    'application/msword',
    'application/vnd.ms-excel',
    'application/vnd.ms-powerpoint',
    'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'application/vnd.openxmlformats-officedocument.presentationml.presentation',
    'application/vnd.ms-excel.sheet.macroEnabled.12',
    'application/vnd.ms-word.document.macroEnabled.12',
    'application/vnd.ms-powerpoint.presentation.macroEnabled.12',
}

OFFICE_EXTENSIONS = {'.doc', '.docx', '.xls', '.xlsx', '.ppt', '.pptx'}
RTL_CHARACTER_RE = re.compile(r'[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF]')
CELL_REF_RE = re.compile(r'([A-Z]+)\d+')


def is_office_document(mime_type: str, filename: str) -> bool:
    if mime_type in OFFICE_MIME_TYPES:
        return True
    return Path(filename).suffix.lower() in OFFICE_EXTENSIONS


def generate_office_preview(source_path: Path) -> Path | None:
    """Generate a PDF preview for an Office document using LibreOffice headless.

    Returns generated preview path on success (PDF or HTML fallback), otherwise None.
    """
    source_path = Path(source_path)
    output_dir = source_path.parent
    output_pdf = output_dir / f'{source_path.stem}.pdf'
    preview_pdf = output_dir / f'{source_path.stem}_preview.pdf'

    try:
        subprocess.run(
            [
                'soffice',
                '--headless',
                '--convert-to',
                'pdf',
                '--outdir',
                str(output_dir),
                str(source_path),
            ],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            timeout=30,
        )
    except (FileNotFoundError, subprocess.SubprocessError, OSError):
        pass
    else:
        if output_pdf.exists():
            try:
                if preview_pdf.exists():
                    preview_pdf.unlink()
                output_pdf.rename(preview_pdf)
            except OSError:
                return output_pdf
            return preview_pdf

    # Fallback path when converter is unavailable:
    # generate a lightweight HTML preview for OOXML documents.
    return _generate_office_html_fallback(source_path)


def _generate_office_html_fallback(source_path: Path) -> Path | None:
    suffix = source_path.suffix.lower()
    if suffix in {'.xlsx', '.xlsm', '.xltx', '.xltm'}:
        return _generate_spreadsheet_html_preview(source_path)
    if suffix == '.docx':
        return _generate_docx_html_preview(source_path)
    if suffix == '.pptx':
        return _generate_pptx_html_preview(source_path)
    return None


def _write_html_preview(source_path: Path, title: str, body_html: str) -> Path:
    preview_path = source_path.parent / f'{source_path.stem}_preview.html'
    document = f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>{escape(title)}</title>
  <style>
    body {{
      margin: 0;
      padding: 16px;
      font-family: "Inter", "Vazirmatn", "Tahoma", "Segoe UI", system-ui, -apple-system, sans-serif;
      background: #f5f7fa;
      color: #1f2937;
    }}
    .card {{
      background: white;
      border: 1px solid #d1d5db;
      border-radius: 8px;
      box-shadow: 0 1px 3px rgba(0,0,0,0.08);
      padding: 12px;
    }}
    h1 {{
      margin: 0 0 12px;
      font-size: 16px;
      line-height: 1.4;
    }}
    table {{
      width: 100%;
      border-collapse: collapse;
      table-layout: fixed;
      font-size: 13px;
    }}
    th, td {{
      border: 1px solid #e5e7eb;
      padding: 6px 8px;
      vertical-align: top;
      word-break: break-word;
      text-align: start;
      unicode-bidi: plaintext;
    }}
    th {{
      background: #f3f4f6;
      text-align: left;
      font-weight: 600;
    }}
    p {{
      margin: 0 0 10px;
      white-space: pre-wrap;
      word-break: break-word;
      line-height: 1.5;
      text-align: start;
      unicode-bidi: plaintext;
    }}
    p[data-dir="rtl"] {{
      direction: rtl;
      font-family: "Vazirmatn", "Tahoma", "Segoe UI", "Inter", system-ui, sans-serif;
    }}
    p[data-dir="ltr"] {{
      direction: ltr;
    }}
    .muted {{
      color: #6b7280;
      font-size: 12px;
      margin-bottom: 10px;
    }}
  </style>
</head>
<body>
  <div class="card">
    <h1>{escape(title)}</h1>
    {body_html}
  </div>
</body>
</html>
"""
    preview_path.write_text(document, encoding='utf-8')
    return preview_path


def _generate_spreadsheet_html_preview(source_path: Path) -> Path | None:
    preview = _generate_spreadsheet_html_preview_openpyxl(source_path)
    if preview:
        return preview
    return _generate_spreadsheet_html_preview_ooxml(source_path)


def _generate_spreadsheet_html_preview_openpyxl(source_path: Path) -> Path | None:
    try:
        from openpyxl import load_workbook
    except Exception:
        return None

    try:
        workbook = load_workbook(source_path, data_only=True, read_only=True)
    except Exception:
        return None

    try:
        sheet = workbook.active
        rows = []
        max_rows = 60
        max_cols = 20
        for row in sheet.iter_rows(min_row=1, max_row=max_rows, max_col=max_cols, values_only=True):
            if all(cell is None for cell in row):
                continue
            cells = ''.join(
                f'<td>{escape(str(cell)) if cell is not None else ""}</td>'
                for cell in row
            )
            rows.append(f'<tr>{cells}</tr>')

        if not rows:
            rows.append('<tr><td class="muted">No visible cells in the preview range.</td></tr>')

        body = (
            f'<div class="muted">Sheet: {escape(sheet.title)} | First {max_rows} rows, {max_cols} columns</div>'
            f'<table>{"".join(rows)}</table>'
        )
        return _write_html_preview(source_path, source_path.name, body)
    except Exception:
        return None
    finally:
        workbook.close()


def _generate_spreadsheet_html_preview_ooxml(source_path: Path) -> Path | None:
    namespace = {
        's': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
        'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    }
    relationship_namespace = {'p': 'http://schemas.openxmlformats.org/package/2006/relationships'}
    max_rows = 60
    max_cols = 20

    try:
        with zipfile.ZipFile(source_path) as archive:
            workbook_xml = archive.read('xl/workbook.xml')
            workbook_root = ET.fromstring(workbook_xml)

            sheet_nodes = workbook_root.findall('.//s:sheets/s:sheet', namespace)
            if not sheet_nodes:
                return None
            first_sheet = sheet_nodes[0]
            sheet_name = first_sheet.attrib.get('name', 'Sheet1')
            relationship_id = first_sheet.attrib.get(f'{{{namespace["r"]}}}id')

            sheet_path = 'xl/worksheets/sheet1.xml'
            if relationship_id:
                rel_xml = archive.read('xl/_rels/workbook.xml.rels')
                rel_root = ET.fromstring(rel_xml)
                for relation in rel_root.findall('.//p:Relationship', relationship_namespace):
                    if relation.attrib.get('Id') != relationship_id:
                        continue
                    target = relation.attrib.get('Target', '')
                    if target.startswith('/'):
                        sheet_path = target.lstrip('/')
                    else:
                        cleaned = target.removeprefix('./')
                        sheet_path = f'xl/{cleaned.lstrip("/")}'
                    break

            if sheet_path not in archive.namelist():
                sheet_candidates = sorted(
                    name for name in archive.namelist()
                    if name.startswith('xl/worksheets/sheet') and name.endswith('.xml')
                )
                if not sheet_candidates:
                    return None
                sheet_path = sheet_candidates[0]

            shared_strings: list[str] = []
            if 'xl/sharedStrings.xml' in archive.namelist():
                shared_xml = archive.read('xl/sharedStrings.xml')
                shared_root = ET.fromstring(shared_xml)
                for si in shared_root.findall('.//s:si', namespace):
                    pieces = []
                    for text_node in si.findall('.//s:t', namespace):
                        if text_node.text:
                            pieces.append(text_node.text)
                    shared_strings.append(''.join(pieces))

            sheet_xml = archive.read(sheet_path)
            sheet_root = ET.fromstring(sheet_xml)
    except (FileNotFoundError, KeyError, OSError, zipfile.BadZipFile, ET.ParseError):
        return None

    rows: list[str] = []
    for row_node in sheet_root.findall('.//s:sheetData/s:row', namespace):
        cells_by_index: dict[int, str] = {}

        for cell_index, cell_node in enumerate(row_node.findall('s:c', namespace), start=1):
            ref = (cell_node.attrib.get('r') or '').upper()
            match = CELL_REF_RE.match(ref)
            column_index = cell_index
            if match:
                column_index = _column_letters_to_index(match.group(1))
            if column_index > max_cols:
                continue

            text = _spreadsheet_cell_text(cell_node, namespace, shared_strings)
            if text:
                cells_by_index[column_index] = escape(text)

        if not cells_by_index:
            continue

        max_used_col = min(max(cells_by_index.keys()), max_cols)
        html_cells = []
        for column_index in range(1, max_used_col + 1):
            html_cells.append(f'<td>{cells_by_index.get(column_index, "")}</td>')
        rows.append(f'<tr>{"".join(html_cells)}</tr>')

        if len(rows) >= max_rows:
            break

    if not rows:
        rows.append('<tr><td class="muted">No visible cells in the preview range.</td></tr>')

    body = (
        f'<div class="muted">Sheet: {escape(sheet_name)} | First {max_rows} rows, {max_cols} columns</div>'
        f'<table>{"".join(rows)}</table>'
    )
    return _write_html_preview(source_path, source_path.name, body)


def _column_letters_to_index(letters: str) -> int:
    index = 0
    for letter in letters:
        index = index * 26 + (ord(letter) - 64)
    return index


def _spreadsheet_cell_text(cell_node: ET.Element, namespace: dict[str, str], shared_strings: list[str]) -> str:
    cell_type = cell_node.attrib.get('t')

    if cell_type == 'inlineStr':
        inline_text = []
        for text_node in cell_node.findall('.//s:is//s:t', namespace):
            if text_node.text:
                inline_text.append(text_node.text)
        return ''.join(inline_text).strip()

    value_node = cell_node.find('s:v', namespace)
    raw_value = (value_node.text or '').strip() if value_node is not None and value_node.text else ''
    if not raw_value:
        return ''

    if cell_type == 's':
        try:
            string_index = int(float(raw_value))
        except ValueError:
            return raw_value
        if 0 <= string_index < len(shared_strings):
            return shared_strings[string_index]
        return raw_value

    return raw_value


def _generate_docx_html_preview(source_path: Path) -> Path | None:
    namespace = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    try:
        with zipfile.ZipFile(source_path) as archive:
            xml_content = archive.read('word/document.xml')
    except (FileNotFoundError, KeyError, OSError, zipfile.BadZipFile):
        return None

    try:
        root = ET.fromstring(xml_content)
    except ET.ParseError:
        return None

    paragraphs = []
    for paragraph in root.findall('.//w:p', namespace):
        text_parts = [node.text for node in paragraph.findall('.//w:t', namespace) if node.text]
        text = ''.join(text_parts).strip()
        if text:
            text_direction = 'rtl' if RTL_CHARACTER_RE.search(text) else 'ltr'
            paragraphs.append(f'<p data-dir="{text_direction}" dir="{text_direction}">{escape(text)}</p>')

    if not paragraphs:
        paragraphs.append('<p class="muted">No readable text found in this document.</p>')

    body = ''.join(paragraphs[:120])
    return _write_html_preview(source_path, source_path.name, body)


def _generate_pptx_html_preview(source_path: Path) -> Path | None:
    try:
        with zipfile.ZipFile(source_path) as archive:
            slide_files = sorted(
                name for name in archive.namelist()
                if name.startswith('ppt/slides/slide') and name.endswith('.xml')
            )
            if not slide_files:
                return None

            slide_sections = []
            for slide_name in slide_files[:30]:
                xml_content = archive.read(slide_name)
                try:
                    root = ET.fromstring(xml_content)
                except ET.ParseError:
                    continue
                texts = []
                for node in root.iter():
                    if node.tag.endswith('}t') and node.text:
                        text = node.text.strip()
                        if text:
                            texts.append(text)
                if texts:
                    slide_title = slide_name.rsplit('/', 1)[-1].replace('.xml', '')
                    joined = escape(' | '.join(texts[:20]))
                    slide_sections.append(f'<p><strong>{escape(slide_title)}:</strong> {joined}</p>')
    except (FileNotFoundError, OSError, zipfile.BadZipFile):
        return None

    if not slide_sections:
        slide_sections = ['<p class="muted">No readable text found in slides.</p>']

    body = ''.join(slide_sections)
    return _write_html_preview(source_path, source_path.name, body)
